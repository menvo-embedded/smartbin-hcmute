"""Sinh file báo cáo tuần 7 (Word tổng kết, Slide tóm tắt) và Chương 1 cuốn báo cáo.
Chạy: python docs/bao-cao/tao_bao_cao.py"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt as PPt
from pptx.dml.color import RGBColor as PRGB

OUT = os.path.dirname(os.path.abspath(__file__))
REPO = 'https://github.com/menvo-embedded/smartbin-hcmute'
TUAN = 'Tuần 7 (22/09 – 28/09/2026)'
GREEN = RGBColor(0x15, 0x80, 0x3D)


def base_doc():
    d = Document()
    for s in d.sections:
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)
        s.left_margin, s.right_margin = Cm(3), Cm(2)
    st = d.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.3
    for lvl, size in ((1, 16), (2, 14), (3, 13)):
        h = d.styles[f'Heading {lvl}']
        h.font.name = 'Times New Roman'
        h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
    return d


def para(d, text, bold=False, italic=False, align=None, size=None):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'justify':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullets(d, items):
    for it in items:
        p = d.add_paragraph(style='List Bullet')
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def table(d, header, rows, widths=None, caption=None):
    if caption:
        para(d, caption, bold=True, align='center', size=12)
    t = d.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(12)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(12)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    d.add_paragraph()
    return t


def cover(d, title, subtitle):
    para(d, 'TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT TP. HỒ CHÍ MINH', bold=True, align='center')
    para(d, 'Môn: Phát triển ứng dụng di động — GVHD: Huỳnh Hoàng Hà', align='center')
    d.add_paragraph()
    para(d, title, bold=True, align='center', size=18)
    para(d, subtitle, bold=True, align='center', size=15)
    para(d, f'Đề tài: SmartBin — Ứng dụng quản lý hệ thống thùng rác thông minh', align='center')
    para(d, f'Mã nguồn: {REPO}', italic=True, align='center', size=11)
    d.add_paragraph()


# ---------------------------------------------------------------- dữ liệu chung
LAM_DUOC = [
    ('Hoàn thiện đủ 4 luồng người dùng trên dữ liệu thật (Supabase): ',
     'kiosk cộng đồng (không cần đăng nhập), hộ gia đình, nhân viên thu gom, quản lý. '
     'Bỏ toàn bộ dữ liệu giả; mọi thao tác ghi đi qua SQLite + hàng đợi đồng bộ (dùng được khi mất mạng).'),
    ('AI nhận diện rác theo bài báo khoa học: ',
     'tích hợp mô hình MobileCLIP-S0 (Apple, CVPR 2024) phân loại zero-shot 3 nhóm hữu cơ / vô cơ / tái chế, '
     'chạy TFLite ngay trên điện thoại. Độ chính xác 93,16% trên 2.749 ảnh thử; trích nguồn IEEE trong ml-research/README.md.'),
    ('Tự nhận diện qua camera, không cần chụp: ',
     'đưa rác ra trước camera → AI nhận ra loại → tự mở đúng ngăn (hộ gia đình + kiosk). '
     'Chạy GPU ~0,5 s/khung; vật nằm yên chỉ tính 1 lần.'),
    ('Hệ thống tự vận hành (chạy trên server): ',
     'thùng ≥60% tự lên lịch thu gom ca gần nhất, ≥80% tạo/nâng việc khẩn; tự giao nhân viên đang trực ít việc nhất; '
     'pg_cron 5 phút/lần tự nhắc việc trễ, dự báo đầy theo tốc độ bỏ rác, báo thùng mất kết nối.'),
    ('Thông báo realtime: ', 'nhân viên/quản lý/hộ gia đình nhận thông báo hệ thống ngay trên điện thoại (Supabase Realtime).'),
    ('Tính năng mới: ', 'dự báo thùng đầy, lộ trình thu gom tối ưu (láng giềng gần nhất + 2-opt), nhân viên chấm chất lượng '
     'phân loại (+5 điểm), tác động môi trường + huy hiệu + chuỗi ngày, tra cứu "rác này bỏ ngăn nào?", ảnh nghiệm thu lưu kho riêng tư.'),
    ('Giao diện: ', 'thanh tab dưới cho cả nhân viên (Công việc / Lộ trình / Cá nhân) và quản lý '
     '(Thùng rác / Điều phối / Tự động / Cá nhân); bản đồ thiết bị.'),
    ('Kiểm thử: ', '44/44 phép kiểm tra server (trigger, phân quyền RLS từng vai trò, tự động hoá), 25/25 kiểm thử logic app, '
     'test thủ công trên điện thoại thật; sửa 8 lỗi phát hiện được (camera mở nắp lặp, lịch sử sai số, bản đồ trống...).'),
    ('Bản cài độc lập: ', 'APK release chạy không cần máy tính (SmartBin-demo.apk).'),
]

PHAN_CONG = [
    ('1', 'Khung dự án, SQLite + hàng đợi đồng bộ, Supabase (schema, RLS)', 'src/core, supabase/schema.sql', 'Hoàn thành'),
    ('2', 'Luồng Hộ gia đình (trang chủ, lịch sử, thống kê, cá nhân, điểm/huy hiệu)', 'src/app/(household), features/stats', 'Hoàn thành'),
    ('3', 'Luồng Kiosk cộng đồng (bỏ rác không đăng nhập, màn chờ, cảm ơn)', 'src/app/(user)', 'Hoàn thành'),
    ('4', 'Luồng Nhân viên thu gom (việc, lộ trình tối ưu, nghiệm thu ảnh)', 'src/app/(collector), features/collection', 'Hoàn thành'),
    ('5', 'Luồng Quản lý (thiết bị, bản đồ, điều phối, lịch ca)', 'src/app/(admin), features/admin', 'Hoàn thành'),
    ('6', 'Tự động hoá phía server + thông báo realtime', 'migrations_manual/2026-09-28_tu_dong_hoa.sql, features/notifications', 'Hoàn thành'),
    ('7', 'AI nhận diện rác (MobileCLIP-S0) + camera trực tiếp', 'src/ml, ml-research', 'Hoàn thành'),
    ('8', 'Phần cứng ESP32 + servo, kết nối BLE thật', 'firmware, core/ble', 'Đang tích hợp'),
    ('9', 'Kiểm thử (server, logic, thủ công trên máy thật)', 'supabase/tests, scripts/test-logic.ts', 'Hoàn thành'),
    ('10', 'Cuốn báo cáo: Chương 1 Tổng quan', 'docs/bao-cao', 'Hoàn thành'),
]

TIEN_DO = [
    ('Tuần 2', 'Khung dự án, hạ tầng (Supabase, SQLite, BLE giả lập), schema + RLS', '100%'),
    ('Tuần 3', 'Đủ màn hình 4 vai trò, luồng bỏ rác đầu-cuối', '100%'),
    ('Tuần 4', 'Giao diện theo thiết kế, thống kê, hồ sơ', '100%'),
    ('Tuần 5', 'Bản đồ thiết bị, repo GitHub, dựng phần cứng, bắt đầu luồng Hộ gia đình', '100%'),
    ('Tuần 6', 'Hoàn thiện Hộ gia đình, đổi sang 3 loại rác, dữ liệu thật', '100%'),
    ('Tuần 7', 'AI MobileCLIP + camera trực tiếp, tự vận hành, thông báo, tab dưới, kiểm thử, Chương 1', '100%'),
    ('Tuần 8', 'Chương 2 Cơ sở lý thuyết; tích hợp BLE với ESP32 thật', 'Dự kiến'),
    ('Tuần 9–10', 'Chương 3 Phân tích & thiết kế hệ thống', 'Dự kiến'),
    ('Tuần 11', 'Chương 4 Phân tích kết quả (benchmark AI, kiểm thử)', 'Dự kiến'),
    ('Tuần 12', 'Chương 5 Kết luận, hoàn thiện & nộp', 'Dự kiến'),
]

TUAN_SAU = [
    'Viết Chương 2 — Cơ sở lý thuyết (React Native/Expo, Supabase, local-first, CLIP/MobileCLIP, TFLite, BLE).',
    'Tích hợp BLE với thùng ESP32 thật (thay MockBinController), demo mở nắp thật.',
    'Đo benchmark AI trên nhiều điện thoại (thời gian/khung, GPU vs CPU) cho Chương 4.',
]


# ---------------------------------------------------------------- 1. Word tổng kết tuần
def tong_ket():
    d = base_doc()
    cover(d, 'BÁO CÁO TIẾN ĐỘ HÀNG TUẦN', TUAN)

    d.add_heading('1. Tuần này làm thêm được gì', level=1)
    bullets(d, LAM_DUOC)

    d.add_heading('2. Demo kết quả', level=1)
    bullets(d, [
        'Cài SmartBin-demo.apk (bản release, chạy độc lập). Tài khoản test mật khẩu 123456: '
        'user1 hộ gia đình, user2/user4 nhân viên, user3 quản lý; kiosk không cần tài khoản.',
        'Kịch bản demo chi tiết: docs/KICH-BAN-DEMO.md (điểm nhấn: hệ thống tự vận hành với 2 máy, '
        'camera AI tự mở nắp, chế độ offline).',
    ])

    d.add_heading('3. Mã nguồn trên GitHub', level=1)
    para(d, REPO, italic=True)
    table(d, ['Thư mục', 'Nội dung'], [
        ('src/app/', 'Màn hình theo vai trò: (household), (user) kiosk, (collector), (admin), (auth)'),
        ('src/features/', 'Nghiệp vụ: sorting, collection, admin, notifications, stats...'),
        ('src/core/', 'Hạ tầng: Supabase, SQLite, hàng đợi đồng bộ, BLE, thông báo'),
        ('src/ml/', 'AI: MobileCLIP-S0 TFLite, camera nhận diện trực tiếp'),
        ('supabase/', 'schema.sql, migration đã chạy, kiểm thử server, dữ liệu demo'),
        ('ml-research/', 'Nghiên cứu mô hình, đánh giá, trích dẫn IEEE'),
        ('docs/', 'Kiến trúc, kịch bản demo, báo cáo tuần, cuốn báo cáo'),
    ], widths=[4, 12])

    d.add_heading('4. Bảng phân công công việc', level=1)
    para(d, 'Ghi chú: điền họ tên – MSSV người phụ trách từng hạng mục.', italic=True, size=11)
    table(d, ['STT', 'Hạng mục', 'Thư mục / file', 'Người phụ trách', 'Trạng thái'],
          [(a, b, c, '', e) for a, b, c, e in PHAN_CONG], widths=[1.2, 6, 4.5, 3, 2.3],
          caption='Bảng 1: Phân công công việc từng thành viên')

    d.add_heading('5. Bảng quản lý tiến độ theo tuần', level=1)
    table(d, ['Tuần', 'Nội dung', 'Tiến độ'], TIEN_DO, widths=[2.2, 11.5, 2.5],
          caption='Bảng 2: Quản lý tiến độ thực tế theo tuần')

    d.add_heading('6. Kế hoạch tuần sau', level=1)
    bullets(d, TUAN_SAU)
    d.save(os.path.join(OUT, 'Tuan07_TongKet.docx'))


# ---------------------------------------------------------------- 2. Slide tóm tắt
def slide():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    blank = p.slide_layouts[6]

    def add(title, lines, sub=None):
        s = p.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, p.slide_width, Inches(1.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PRGB(0x16, 0xA3, 0x4A)
        bar.line.fill.background()
        tf = bar.text_frame
        tf.text = title
        tf.paragraphs[0].runs[0].font.size = PPt(30)
        tf.paragraphs[0].runs[0].font.bold = True
        tf.paragraphs[0].runs[0].font.color.rgb = PRGB(255, 255, 255)
        box = s.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12.1), Inches(5.8))
        t = box.text_frame
        t.word_wrap = True
        if sub:
            t.text = sub
            t.paragraphs[0].runs[0].font.size = PPt(18)
            t.paragraphs[0].runs[0].font.italic = True
        for i, line in enumerate(lines):
            para_ = t.paragraphs[0] if (i == 0 and not sub) else t.add_paragraph()
            if isinstance(line, tuple):
                r = para_.add_run()
                r.text = '• ' + line[0]
                r.font.bold = True
                r.font.size = PPt(18)
                r2 = para_.add_run()
                r2.text = line[1]
                r2.font.size = PPt(16)
            else:
                r = para_.add_run()
                r.text = '• ' + line
                r.font.size = PPt(18)
            para_.space_after = PPt(8)
        return s

    add('SmartBin — Báo cáo ' + TUAN, [
        'Môn Phát triển ứng dụng di động — GVHD: Huỳnh Hoàng Hà',
        'Đề tài: Ứng dụng quản lý hệ thống thùng rác thông minh',
        'Mã nguồn: ' + REPO,
    ])
    add('1. Tuần này làm thêm được gì', [(a, '') for a, _ in LAM_DUOC[:5]])
    add('1. Tuần này làm thêm được gì (tiếp)', [(a, '') for a, _ in LAM_DUOC[5:]])
    add('2. Hệ thống tự vận hành', [
        'Thùng ≥60% → tự lên lịch ca gần nhất; ≥80% → việc khẩn',
        'Việc mới → tự giao nhân viên đang trực ít việc nhất',
        'Mỗi 5 phút: nhắc việc trễ, dự báo đầy, phát hiện thùng mất kết nối',
        'Mọi sự kiện → thông báo realtime đúng người (nhân viên / quản lý / hộ gia đình)',
        'Quản lý bật/tắt quy tắc, chọn nhân viên đang trực trong tab "Tự động"',
    ])
    add('3. AI nhận diện rác — MobileCLIP-S0 (CVPR 2024)', [
        'Zero-shot: so ảnh với 180 câu mô tả, không cần tự huấn luyện',
        'Độ chính xác 93,16% (2.749 ảnh), chạy TFLite trên điện thoại, không cần mạng',
        'Camera trực tiếp: đưa rác ra trước camera → tự mở đúng ngăn (~0,5 s/khung GPU)',
        'Trích nguồn IEEE: ml-research/README.md',
    ])
    add('4. Kiểm thử & sửa lỗi', [
        'Server: 44/44 (trigger, RLS từng vai trò, tự động hoá)',
        'Logic app: 25/25 (dự báo, lộ trình, tác động môi trường, tra cứu, AI)',
        'Test trên điện thoại thật cả 4 vai trò + offline + realtime 2 máy',
        'Sửa 8 lỗi: camera mở nắp lặp, lịch sử sai số, bản đồ trống, kiosk kẹt màn cảm ơn...',
    ])
    add('5. Phân công & tiến độ', [f'{a}. {b} — {e}' for a, b, c, e in PHAN_CONG],
        sub='Chi tiết người phụ trách: Bảng 1 trong file Word tổng kết')
    add('6. Kế hoạch tuần sau', TUAN_SAU)
    p.save(os.path.join(OUT, 'Tuan07_TomTat.pptx'))


# ---------------------------------------------------------------- 3. Chương 1
def chuong1():
    d = base_doc()
    viet_chuong1(d)
    d.save(os.path.join(OUT, 'BaoCao_Chuong1_TongQuan.docx'))


def viet_chuong1(d):
    h = d.add_heading('CHƯƠNG 1: TỔNG QUAN', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading('1.1. Đặt vấn đề', level=2)
    for t in [
        'Rác thải sinh hoạt là một trong những vấn đề môi trường lớn của các đô thị Việt Nam. Theo Luật Bảo vệ môi trường 2020, '
        'từ ngày 31/12/2024 các hộ gia đình, cá nhân bắt buộc phải phân loại chất thải rắn sinh hoạt tại nguồn thành ba nhóm: '
        'chất thải có khả năng tái sử dụng, tái chế; chất thải thực phẩm (hữu cơ); và chất thải rắn sinh hoạt khác. '
        'Tuy nhiên, trên thực tế người dân vẫn thường nhầm lẫn khi phân loại, còn việc thu gom chủ yếu theo lịch cố định: '
        'có thùng đầy tràn chưa được thu gom trong khi có thùng còn trống vẫn được ghé qua, gây lãng phí nhân lực và mất vệ sinh.',
        'Sự phát triển của điện thoại thông minh, Internet vạn vật (IoT) và trí tuệ nhân tạo (AI) chạy trực tiếp trên thiết bị '
        'mở ra hướng giải quyết: thùng rác có thể tự mở đúng ngăn khi nhận ra loại rác, tự báo mức đầy, và hệ thống có thể tự lên '
        'lịch, tự phân công nhân viên thu gom dựa trên dữ liệu thực tế thay vì lịch cố định.',
        'Xuất phát từ thực tế đó, nhóm chọn đề tài "SmartBin — Ứng dụng quản lý hệ thống thùng rác thông minh", xây dựng một '
        'ứng dụng di động kết nối người dân, nhân viên thu gom và người quản lý trên cùng một nền tảng dữ liệu thời gian thực.',
    ]:
        para(d, t, align='justify')

    d.add_heading('1.2. Mục tiêu của đề tài', level=2)
    para(d, 'Mục tiêu tổng quát: xây dựng ứng dụng di động hoàn chỉnh hỗ trợ phân loại rác tại nguồn và quản lý thu gom rác '
            'một cách tự động. Các mục tiêu cụ thể:', align='justify')
    bullets(d, [
        'Hỗ trợ người dân phân loại đúng 3 nhóm rác (hữu cơ, vô cơ, tái chế): chọn tay, tra cứu, hoặc để AI nhận diện qua camera '
        'và tự mở đúng ngăn thùng.',
        'Khuyến khích phân loại bằng điểm thưởng, huy hiệu, thống kê tác động môi trường (khối lượng rác, lượng CO₂ giảm được).',
        'Giúp nhân viên thu gom nhận việc, xem lộ trình tối ưu, chụp ảnh nghiệm thu và đánh giá chất lượng phân loại.',
        'Giúp người quản lý giám sát mức đầy các thùng trên bản đồ theo thời gian thực, điều phối nhân viên, lên lịch ca trực.',
        'Tự động hoá vận hành: tự lên lịch khi thùng sắp đầy, tự giao việc, tự nhắc việc trễ, tự phát hiện thiết bị mất kết nối.',
        'Ứng dụng vẫn hoạt động khi mất mạng (ghi cục bộ, đồng bộ lại khi có mạng).',
    ])

    d.add_heading('1.3. Đối tượng và phạm vi nghiên cứu', level=2)
    para(d, 'Đối tượng sử dụng gồm bốn nhóm:', align='justify')
    table(d, ['Đối tượng', 'Chức năng chính'], [
        ('Người dân (kiosk cộng đồng)', 'Bỏ rác tại thùng công cộng, không cần tài khoản; AI tự nhận diện và mở ngăn'),
        ('Hộ gia đình', 'Thùng rác riêng; tích điểm, lịch sử, thống kê, huy hiệu, lịch thu gom'),
        ('Nhân viên thu gom', 'Nhận việc, lộ trình tối ưu, nghiệm thu bằng ảnh, chấm chất lượng phân loại'),
        ('Người quản lý', 'Giám sát thùng, bản đồ, điều phối, lịch ca, cấu hình tự động hoá'),
    ], widths=[5, 11], caption='Bảng 1.1: Đối tượng sử dụng và chức năng')
    para(d, 'Phạm vi:', bold=True)
    bullets(d, [
        'Ứng dụng di động Android (React Native / Expo), máy chủ Supabase (PostgreSQL, xác thực, lưu trữ ảnh, realtime).',
        'Phân loại theo 3 nhóm rác đúng quy định: hữu cơ, vô cơ, tái chế.',
        'Phần cứng ở mức mô hình: một thùng rác 3 ngăn điều khiển bằng ESP32 + servo qua Bluetooth Low Energy (BLE); '
        'có chế độ giả lập để demo khi không có phần cứng.',
        'Mô hình AI sử dụng mô hình đã công bố (MobileCLIP-S0, CVPR 2024), chạy trên điện thoại; không huấn luyện mô hình mới.',
    ])

    d.add_heading('1.4. Phương pháp thực hiện', level=2)
    bullets(d, [
        ('Khảo sát: ', 'tìm hiểu quy định phân loại rác tại nguồn, các nghiên cứu về phân loại rác bằng học sâu và mô hình '
         'thị giác – ngôn ngữ (CLIP, MobileCLIP).'),
        ('Thiết kế: ', 'kiến trúc tách lớp hạ tầng (core) và nghiệp vụ (features); mô hình dữ liệu PostgreSQL với phân quyền '
         'theo vai trò (Row Level Security); thiết kế "ghi cục bộ trước, đồng bộ sau" (local-first) với SQLite và hàng đợi đồng bộ.'),
        ('Hiện thực: ', 'React Native + Expo Router, Zustand, TanStack Query; Supabase (trigger, pg_cron, Realtime); '
         'TensorFlow Lite + VisionCamera cho AI; BLE cho phần cứng.'),
        ('Kiểm thử: ', 'kiểm thử tự động phía máy chủ (44 phép kiểm tra trigger, phân quyền, tự động hoá), kiểm thử logic ứng '
         'dụng (25 phép), đánh giá AI trên 2.749 ảnh, kiểm thử thủ công trên điện thoại thật.'),
        ('Quản lý dự án: ', 'mã nguồn trên GitHub, báo cáo tiến độ hằng tuần, bảng phân công công việc.'),
    ])

    d.add_heading('1.5. Kết quả dự kiến', level=2)
    bullets(d, [
        'Ứng dụng Android hoàn chỉnh cho 4 nhóm người dùng, chạy trên dữ liệu thật.',
        'Chức năng AI nhận diện rác trên thiết bị với độ chính xác trên 90%, tự mở đúng ngăn thùng.',
        'Hệ thống thu gom tự vận hành: tự lên lịch, tự phân công, tự cảnh báo.',
        'Mô hình thùng rác thông minh ESP32 kết nối với ứng dụng.',
        'Cuốn báo cáo, mã nguồn và kịch bản demo.',
    ])

    d.add_heading('1.6. Bố cục báo cáo', level=2)
    bullets(d, [
        ('Chương 1 – Tổng quan: ', 'đặt vấn đề, mục tiêu, đối tượng, phạm vi, phương pháp.'),
        ('Chương 2 – Cơ sở lý thuyết: ', 'các công nghệ và mô hình được sử dụng.'),
        ('Chương 3 – Phân tích và thiết kế hệ thống: ', 'yêu cầu, kiến trúc, cơ sở dữ liệu, thiết kế giao diện và các luồng xử lý.'),
        ('Chương 4 – Phân tích kết quả: ', 'kết quả hiện thực, kiểm thử, đánh giá AI.'),
        ('Chương 5 – Kết luận: ', 'kết quả đạt được, hạn chế, hướng phát triển.'),
    ])


if __name__ == '__main__':
    tong_ket()
    slide()
    chuong1()
    print('Đã tạo:', os.listdir(OUT))
