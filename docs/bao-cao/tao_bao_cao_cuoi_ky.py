"""Sinh cuốn báo cáo cuối kỳ SmartBin: Chương 1–5 (mỗi chương một file) + cuốn đầy đủ.
Chạy: python docs/bao-cao/tao_bao_cao_cuoi_ky.py"""
import os
from PIL import Image
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch

from tao_bao_cao import OUT, REPO, base_doc, para, bullets, table, viet_chuong1

HINH = os.path.join(OUT, 'hinh')
J = 'justify'


def text(d, *paras):
    for t in paras:
        para(d, t, align=J)


def code(d, t):
    p = d.add_paragraph()
    r = p.add_run(t)
    r.font.name = 'Consolas'
    r.font.size = Pt(10.5)
    p.paragraph_format.left_indent = Cm(0.8)
    return p


def figure(d, path, caption, width_cm=15):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    para(d, caption, italic=True, align='center', size=12)


def phone_shot(name):
    """Cắt thanh trạng thái / thanh điều hướng của ảnh chụp màn hình điện thoại."""
    src = os.path.join(HINH, name + '.png')
    dst = os.path.join(HINH, name + '_crop.png')
    if not os.path.exists(dst):
        im = Image.open(src)
        w, h = im.size
        im.crop((0, 100, w, h - 165)).save(dst)
    return dst


def screens(d, items, caption):
    """Các ảnh màn hình đặt cạnh nhau (tối đa 3 ảnh/hàng)."""
    t = d.add_table(rows=2, cols=len(items))
    for i, (name, label) in enumerate(items):
        c = t.rows[0].cells[i]
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].add_run().add_picture(phone_shot(name), width=Cm(4.8))
        lc = t.rows[1].cells[i]
        lc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = lc.paragraphs[0].add_run(label)
        r.font.size = Pt(11)
        r.italic = True
    para(d, caption, italic=True, align='center', size=12)


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- sơ đồ kiến trúc
def ve_kien_truc():
    path = os.path.join(HINH, 'kien_truc.png')
    fig, ax = plt.subplots(figsize=(12, 7.2), dpi=150)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 7.2)
    ax.axis('off')

    def box(x, y, w, h, title, lines, color):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.05,rounding_size=0.15',
                                    fc=color, ec='#15803d', lw=1.5))
        ax.text(x + w / 2, y + h - 0.28, title, ha='center', va='top', fontsize=11.5, weight='bold')
        ax.text(x + w / 2, y + h - 0.72, '\n'.join(lines), ha='center', va='top', fontsize=9.2, linespacing=1.5)

    def arrow(x1, y1, x2, y2, label='', both=False):
        ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle='<->' if both else '->', lw=1.6, color='#374151'))
        if label:
            ax.text((x1 + x2) / 2, (y1 + y2) / 2 + 0.12, label, ha='center', fontsize=8.8, color='#374151',
                    bbox=dict(fc='white', ec='none', pad=1))

    box(0.2, 3.1, 4.3, 3.9, 'Ứng dụng di động (React Native / Expo)', [
        'Màn hình theo vai trò (Expo Router, tab dưới)',
        'Hộ gia đình · Kiosk · Nhân viên · Quản lý',
        '─────────────',
        'features/: nghiệp vụ (bỏ rác, thu gom, điều phối)',
        'core/: SQLite + hàng đợi đồng bộ, Supabase,',
        'BLE (thật / giả lập), thông báo cục bộ',
        'ml/: MobileCLIP-S0 TFLite + camera',
        'Zustand · TanStack Query'], '#ecfdf5')
    box(7.3, 3.1, 4.5, 3.9, 'Supabase (máy chủ)', [
        'PostgreSQL: 7 bảng + RLS theo vai trò',
        'Trigger: cộng điểm, mức đầy, tạo việc,',
        'tự giao việc, sinh thông báo',
        'pg_cron: run_automation() mỗi 5 phút',
        'Auth (JWT) · Storage (ảnh nghiệm thu)',
        'Realtime (WebSocket)'], '#eff6ff')
    box(0.2, 0.2, 4.3, 2.2, 'Thùng rác thông minh', [
        'ESP32 · GATT server BLE',
        '3 servo (hữu cơ / vô cơ / tái chế)',
        'Lệnh JSON {"cmd":"open","bin":...}'], '#fefce8')
    box(7.3, 0.2, 4.5, 2.2, 'Mô hình AI (xử lý ngoại tuyến trên PC)', [
        'MobileCLIP-S0 (Apple, CVPR 2024)',
        'ONNX → TFLite (onnx2tf), lượng tử hoá',
        '180 câu mô tả → vector văn bản (JSON)'], '#fdf2f8')

    arrow(4.6, 5.6, 7.2, 5.6, 'HTTPS (REST) — đồng bộ hàng đợi', both=True)
    arrow(7.2, 4.4, 4.6, 4.4, 'Realtime: thay đổi dữ liệu + thông báo')
    arrow(2.35, 3.0, 2.35, 2.5, '', both=True)
    ax.text(2.5, 2.72, 'BLE', fontsize=9.5, color='#374151')
    arrow(7.2, 1.3, 4.6, 3.3, 'mô hình .tflite + text_embeddings.json đóng gói vào app')
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


# ---------------------------------------------------------------- CHƯƠNG 2
def viet_chuong2(d):
    h = d.add_heading('CHƯƠNG 2: CƠ SỞ LÝ THUYẾT', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading('2.1. Phân loại chất thải rắn sinh hoạt tại nguồn', level=2)
    text(d,
         'Luật Bảo vệ môi trường số 72/2020/QH14 (Điều 75) và Nghị định 08/2022/NĐ-CP quy định chất thải rắn sinh hoạt '
         'phải được phân loại tại nguồn thành ba nhóm: (1) chất thải có khả năng tái sử dụng, tái chế; (2) chất thải thực '
         'phẩm; (3) chất thải rắn sinh hoạt khác. Ứng dụng SmartBin dùng đúng ba nhóm này với tên gọi quen thuộc: '
         'tái chế, hữu cơ và vô cơ. Việc phân loại đúng giúp tăng tỷ lệ tái chế, giảm lượng rác chôn lấp và giảm phát thải '
         'khí mê-tan từ rác hữu cơ.')

    d.add_heading('2.2. React Native và Expo', level=2)
    text(d,
         'React Native là framework của Meta cho phép viết ứng dụng di động bằng JavaScript/TypeScript với thư viện React; '
         'giao diện được hiển thị bằng thành phần gốc (native) của Android/iOS chứ không phải trang web. Phiên bản 0.76 dùng '
         'kiến trúc mới (Fabric, TurboModules) và máy ảo JavaScript Hermes được tối ưu cho di động.',
         'Expo (SDK 52) cung cấp bộ công cụ build, các module dựng sẵn (camera, SQLite, thông báo, vị trí, lưu trữ bảo mật) '
         'và Expo Router — điều hướng theo tệp: mỗi tệp trong thư mục app/ là một màn hình, thư mục có ngoặc như (admin) '
         'là một nhóm màn hình dùng chung một layout (ví dụ thanh tab dưới). Vì app dùng module gốc (TFLite, VisionCamera, BLE) '
         'nên phải dùng development build / bản release thay vì Expo Go.')

    d.add_heading('2.3. Quản lý trạng thái và dữ liệu phía ứng dụng', level=2)
    bullets(d, [
        ('Zustand: ', 'thư viện quản lý trạng thái toàn cục gọn nhẹ (hook create), dùng cho phiên đăng nhập, cấu hình kiosk, '
         'thời điểm cập nhật realtime.'),
        ('TanStack Query: ', 'quản lý dữ liệu lấy từ máy chủ: bộ nhớ đệm theo khoá (queryKey), tự tải lại, làm mới khi dữ liệu '
         'thay đổi (invalidateQueries). Các màn hình dùng chung một khoá sẽ dùng chung một bản dữ liệu.'),
        ('SQLite (expo-sqlite): ', 'cơ sở dữ liệu cục bộ trên điện thoại, lưu sự kiện bỏ rác và hàng đợi đồng bộ.'),
    ])

    d.add_heading('2.4. Supabase', level=2)
    text(d, 'Supabase là nền tảng backend mã nguồn mở xây dựng trên PostgreSQL, cung cấp:')
    bullets(d, [
        ('PostgreSQL: ', 'hệ quản trị CSDL quan hệ; hỗ trợ hàm PL/pgSQL, trigger (tự chạy khi INSERT/UPDATE), ràng buộc CHECK.'),
        ('Auth: ', 'đăng nhập bằng email/mật khẩu, cấp JSON Web Token (JWT); trong SQL, auth.uid() trả về người đang gọi.'),
        ('Row Level Security (RLS): ', 'chính sách phân quyền theo từng dòng dữ liệu. Ví dụ nhân viên chỉ đọc được việc giao cho '
         'mình; kiosk (vai trò anon) chỉ được ghi sự kiện không gắn người dùng. Phân quyền nằm ở CSDL nên không thể vượt qua từ app.'),
        ('Realtime: ', 'đọc nhật ký thay đổi của PostgreSQL (logical replication) và đẩy về app qua WebSocket; RLS vẫn áp dụng.'),
        ('pg_cron: ', 'lập lịch chạy câu lệnh SQL định kỳ (cú pháp cron) ngay trong CSDL.'),
        ('Storage: ', 'lưu tệp (ảnh nghiệm thu) trong bucket riêng tư, truy cập qua đường dẫn có chữ ký, có thời hạn.'),
    ])

    d.add_heading('2.5. Kiến trúc "ghi cục bộ trước" (local-first) và đồng bộ ngoại tuyến', level=2)
    text(d,
         'Thùng rác đặt ngoài trời có thể mất mạng. Theo hướng local-first, mọi thao tác ghi được lưu ngay vào SQLite trên máy '
         'rồi đưa vào một hàng đợi; một bộ máy đồng bộ gửi dần hàng đợi lên máy chủ khi có mạng. Ba kỹ thuật chính:',)
    bullets(d, [
        ('Khoá chống trùng (idempotency): ', 'mỗi sự kiện có local_id (UUID sinh trên máy, cột UNIQUE trên máy chủ), gửi lại '
         'nhiều lần cũng không tạo bản ghi trùng.'),
        ('Thử lại có giới hạn: ', 'lỗi mạng thì dừng và giữ nguyên hàng đợi; lỗi từ máy chủ thì tính một lần thất bại, tối đa 5 lần '
         'để một thao tác hỏng không chặn các thao tác phía sau.'),
        ('Giải quyết xung đột: ', 'last-write-wins theo thời điểm tạo — phù hợp vì dữ liệu sự kiện chủ yếu chỉ ghi thêm.'),
    ])

    d.add_heading('2.6. Bluetooth Low Energy và vi điều khiển ESP32', level=2)
    text(d,
         'BLE là chuẩn Bluetooth tiết kiệm năng lượng. Thiết bị ngoại vi (ESP32) đóng vai trò GATT server, cung cấp các '
         'service gồm nhiều characteristic; điện thoại (GATT client) quét, kết nối, ghi và đọc characteristic. ESP32 là vi '
         'điều khiển Wi-Fi/BLE giá rẻ; servo được điều khiển bằng xung PWM để quay nắp đến góc mong muốn.')

    d.add_heading('2.7. Học sâu cho phân loại ảnh và mô hình thị giác – ngôn ngữ', level=2)
    text(d,
         'Các nghiên cứu phân loại rác thường huấn luyện mạng tích chập (CNN) trên tập ảnh có nhãn, ví dụ TrashNet [3] (6 loại '
         'vật liệu) hay MobileNetV2 [4]. Hạn chế: cần dữ liệu lớn cho đúng các nhóm của bài toán, và muốn đổi nhóm phải huấn luyện lại.',
         'CLIP [2] huấn luyện đồng thời bộ mã hoá ảnh và bộ mã hoá văn bản trên hàng trăm triệu cặp ảnh – chú thích bằng học '
         'tương phản (contrastive learning): ảnh và câu mô tả đúng của nó được kéo lại gần nhau trong cùng một không gian vector. '
         'Nhờ vậy có thể phân loại zero-shot: không cần huấn luyện thêm, chỉ cần so ảnh với các câu mô tả từng lớp.')
    para(d, 'Với vector ảnh v và vector câu mô tả tᵢ đã chuẩn hoá về độ dài 1, xác suất của câu i là:', align=J)
    code(d, 'p(i) = exp(100 · cos(v, tᵢ)) / Σⱼ exp(100 · cos(v, tⱼ))')
    text(d,
         'trong đó 100 là hệ số nhiệt độ (logit scale) học được của CLIP. Kỹ thuật prompt ensembling [2] dùng nhiều câu mô tả '
         'cho mỗi lớp (nhiều đồ vật × nhiều mẫu câu); điểm của một lớp là tổng xác suất các câu thuộc lớp đó.',
         'MobileCLIP [1] (Apple, CVPR 2024) là họ mô hình CLIP thiết kế cho thiết bị di động, huấn luyện theo phương pháp '
         '"multi-modal reinforced training" để giữ độ chính xác cao với độ trễ thấp. Bản nhỏ nhất MobileCLIP-S0 có bộ mã hoá ảnh '
         'lai CNN–Transformer, đầu vào 256×256, đầu ra vector 512 chiều. Hướng zero-shot bằng mô hình thị giác – ngôn ngữ cũng '
         'đã được áp dụng cho phân loại rác nhựa trong [7].')

    d.add_heading('2.8. TensorFlow Lite và lượng tử hoá', level=2)
    text(d,
         'TensorFlow Lite (LiteRT) là môi trường chạy mô hình học sâu trên thiết bị. Delegate cho phép chuyển tính toán sang '
         'GPU (OpenCL/OpenGL) để tăng tốc. Lượng tử hoá giảm độ chính xác số của trọng số: FP16 (16 bit) giảm một nửa dung lượng; '
         'INT8 dynamic range lưu trọng số 8 bit và giải lượng tử khi chạy; lượng tử hoá nguyên hoàn toàn [6] tính toán hoàn toàn '
         'bằng số nguyên nên nhanh nhất trên phần cứng di động.')

    d.add_heading('2.9. Bài toán lộ trình và dự báo', level=2)
    bullets(d, [
        ('Khoảng cách haversine: ', 'khoảng cách trên mặt cầu giữa hai toạ độ (vĩ độ, kinh độ), bán kính Trái Đất 6.371 km.'),
        ('Bài toán người du lịch (TSP): ', 'tìm thứ tự ghé thăm các điểm có tổng quãng đường nhỏ nhất — bài toán NP-khó. Với vài '
         'chục điểm, heuristic láng giềng gần nhất (luôn đi tới điểm gần nhất chưa ghé) kết hợp cải thiện 2-opt [8] (đảo ngược một '
         'đoạn đường nếu làm tổng quãng đường ngắn lại) cho lời giải gần tối ưu trong thời gian rất ngắn.'),
        ('Dự báo tuyến tính: ', 'tốc độ đầy của một ngăn = số lượt bỏ rác trong cửa sổ quan sát × mức đầy mỗi lượt; thời gian còn '
         'lại đến ngưỡng = (ngưỡng − mức hiện tại) / tốc độ.'),
    ])


# ---------------------------------------------------------------- CHƯƠNG 3
def viet_chuong3(d, arch_png):
    h = d.add_heading('CHƯƠNG 3: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading('3.1. Phân tích yêu cầu', level=2)
    d.add_heading('3.1.1. Tác nhân', level=3)
    table(d, ['Tác nhân', 'Mô tả'], [
        ('Người dân (kiosk)', 'Bỏ rác ở thùng công cộng qua màn hình kiosk gắn trên thùng; không cần tài khoản.'),
        ('Hộ gia đình', 'Sở hữu một thùng rác riêng; có tài khoản, tích điểm.'),
        ('Nhân viên thu gom', 'Nhận và thực hiện việc thu gom, nghiệm thu bằng ảnh.'),
        ('Quản lý', 'Giám sát toàn hệ thống, điều phối nhân viên, cấu hình tự động hoá.'),
        ('Hệ thống (máy chủ)', 'Tác nhân tự động: trigger và lịch pg_cron tự lên lịch, giao việc, cảnh báo.'),
        ('Thùng rác ESP32', 'Thiết bị mở ngăn theo lệnh BLE.'),
    ], widths=[4.5, 11.5], caption='Bảng 3.1: Các tác nhân của hệ thống')

    d.add_heading('3.1.2. Yêu cầu chức năng', level=3)
    table(d, ['Mã', 'Tác nhân', 'Chức năng'], [
        ('F01', 'Tất cả', 'Chọn chế độ Cộng đồng / Hộ gia đình; đăng nhập đúng vai trò của chế độ'),
        ('F02', 'Kiosk, Hộ GĐ', 'Bỏ rác: chọn loại, tra cứu, chụp ảnh AI, hoặc camera AI tự nhận diện → mở đúng ngăn'),
        ('F03', 'Kiosk, Hộ GĐ', 'Hoạt động khi mất mạng, tự đồng bộ khi có mạng'),
        ('F04', 'Hộ GĐ', 'Điểm thưởng, lịch sử, thống kê, tác động môi trường, huy hiệu, chuỗi ngày'),
        ('F05', 'Hộ GĐ', 'Xem lịch thu gom kế tiếp; nhận thông báo lên lịch, đã thu gom, đánh giá phân loại'),
        ('F06', 'Nhân viên', 'Danh sách việc, nhận việc, lộ trình tối ưu, chỉ đường'),
        ('F07', 'Nhân viên', 'Hoàn tất việc kèm ảnh nghiệm thu; chấm chất lượng phân loại thùng hộ gia đình'),
        ('F08', 'Quản lý', 'Bản đồ và mức đầy từng ngăn theo thời gian thực; dự báo thời điểm đầy'),
        ('F09', 'Quản lý', 'Giao / đổi / huỷ giao việc, lên lịch ca trực, xem ảnh nghiệm thu'),
        ('F10', 'Quản lý', 'Bật/tắt quy tắc tự động, chọn nhân viên đang trực, xem nhật ký tự động'),
        ('F11', 'Hệ thống', 'Tự lên lịch khi ngăn ≥60%, tạo/nâng việc khẩn khi ≥80%, tự giao việc'),
        ('F12', 'Hệ thống', 'Định kỳ: nhắc/leo thang việc trễ, dự báo theo tốc độ, phát hiện thùng mất kết nối'),
        ('F13', 'Hệ thống', 'Gửi thông báo đúng người, hiển thị thông báo trên điện thoại'),
    ], widths=[1.3, 3, 11.7], caption='Bảng 3.2: Yêu cầu chức năng')

    d.add_heading('3.1.3. Yêu cầu phi chức năng', level=3)
    bullets(d, [
        ('Khả dụng ngoại tuyến: ', 'thao tác bỏ rác không phụ thuộc mạng.'),
        ('Bảo mật: ', 'phân quyền ở CSDL (RLS); điểm thưởng chỉ do máy chủ cộng; ảnh nghiệm thu không công khai.'),
        ('Hiệu năng: ', 'AI chạy trên máy, không gửi ảnh ra ngoài; nhận diện trực tiếp dưới 1 giây/khung.'),
        ('Thời gian thực: ', 'thay đổi hiển thị trên các máy khác trong khoảng 1–2 giây.'),
        ('Dễ bảo trì: ', 'tách hạ tầng và nghiệp vụ; module AI độc lập; có bản giả lập phần cứng để demo.'),
    ])

    d.add_heading('3.2. Kiến trúc hệ thống', level=2)
    figure(d, arch_png, 'Hình 3.1: Kiến trúc tổng thể hệ thống SmartBin')
    text(d,
         'Ứng dụng chia thành các lớp: app/ (màn hình theo vai trò), features/ (nghiệp vụ: bỏ rác, thu gom, điều phối, thông '
         'báo, thống kê), core/ (hạ tầng: Supabase, SQLite, đồng bộ, BLE, thông báo — không biết gì về rác), ml/ (mô hình AI, '
         'nơi duy nhất dùng TFLite và camera) và shared/ (giao diện, kiểu dữ liệu dùng chung). Các feature chỉ nhận kết quả dạng '
         '{loại rác, độ tin cậy}, không phân biệt do người chọn hay AI suy ra.')

    d.add_heading('3.3. Thiết kế cơ sở dữ liệu', level=2)
    table(d, ['Bảng', 'Cột chính', 'Ý nghĩa'], [
        ('profiles', 'id, full_name, role, points, on_duty', 'Người dùng; vai trò household/collector/admin; nhân viên đang trực'),
        ('devices', 'id, code, name, area, latitude, longitude, is_online, last_seen_at, owner_id', 'Thùng rác; owner_id ≠ null là thùng hộ gia đình'),
        ('bins', 'id, device_id, waste_type, fill_level', 'Ba ngăn của mỗi thùng, mức đầy 0..1'),
        ('sort_events', 'id, local_id (unique), device_id, user_id, waste_type, source (manual/ai), confidence', 'Mỗi lượt bỏ rác'),
        ('collection_tasks', 'id, device_id, assignee_id, status, priority, shift, scheduled_date, origin, auto_assigned, '
         'proof_photo_url, sorting_quality, escalated_at', 'Việc thu gom; origin cho biết việc do quản lý hay hệ thống tạo'),
        ('notifications', 'id, recipient_id (null = mọi quản lý), kind, title, body, device_id, task_id, read_at', 'Thông báo / nhật ký tự động'),
        ('automation_settings', 'auto_dispatch, predictive_schedule, escalate_after_min, offline_after_min', 'Cấu hình tự động hoá (1 dòng)'),
    ], widths=[3.3, 7.7, 5], caption='Bảng 3.3: Các bảng dữ liệu')
    text(d, 'Phân quyền (Row Level Security) tiêu biểu:')
    table(d, ['Bảng', 'Chính sách'], [
        ('sort_events', 'Người đăng nhập chỉ ghi sự kiện của chính mình; kiosk (anon) chỉ ghi sự kiện user_id rỗng'),
        ('collection_tasks', 'Nhân viên xem việc giao cho mình và việc chưa ai nhận; hộ gia đình xem việc của thùng nhà mình; '
                             'quản lý xem tất cả'),
        ('notifications', 'Mỗi người chỉ đọc thông báo gửi cho mình; nhật ký chung (recipient null) chỉ quản lý đọc được'),
        ('automation_settings, bins', 'Chỉ quản lý được sửa'),
    ], widths=[4.3, 11.7], caption='Bảng 3.4: Chính sách phân quyền tiêu biểu')

    d.add_heading('3.4. Thiết kế xử lý', level=2)
    d.add_heading('3.4.1. Luồng bỏ rác', level=3)
    bullets(d, [
        'Xác định loại rác (chọn tay / tra cứu / AI).',
        'Gửi lệnh mở ngăn tới thùng qua BLE (hoặc bản giả lập).',
        'Ghi sự kiện vào SQLite với local_id mới → đưa vào hàng đợi đồng bộ → thử gửi ngay nếu có mạng.',
        'Trên máy chủ, trigger khi có sort_events mới: cộng 1 điểm (hộ gia đình), tăng 2% mức đầy ngăn tương ứng, cập nhật '
        'nhịp tim thiết bị; nếu mức đầy vượt ngưỡng thì kích hoạt chuỗi tự động (3.4.4).',
    ])

    d.add_heading('3.4.2. Nhận diện rác trực tiếp bằng camera', level=3)
    text(d, 'Quy trình mỗi khung hình (chạy 2 khung/giây trong frame processor của VisionCamera):')
    code(d, 'Khung camera → cắt giữa, thu về 256×256 RGB → MobileCLIP-S0 (TFLite, GPU) → vector 512 chiều\n'
            '→ cosine với 180 câu mô tả (4 nhóm: hữu cơ, vô cơ, tái chế, "không có rác") → điểm từng nhóm')
    text(d, 'Để tránh mở nhầm, ứng dụng dùng một máy trạng thái:')
    table(d, ['Trạng thái', 'Điều kiện chuyển'], [
        ('Chờ (waiting)', 'Nhóm cao nhất là một loại rác với điểm ≥ 0,8 → Xác nhận'),
        ('Xác nhận (detecting)', '2 khung liên tiếp cùng loại → mở ngăn, gửi sự kiện → Nghỉ; khác loại/không có rác → Chờ'),
        ('Nghỉ (cooldown)', 'Sau 5 giây → Chờ vật rời khung'),
        ('Chờ vật rời khung (clear)', '2 khung liên tiếp không thấy rác → Chờ (vật nằm yên chỉ tính một lần)'),
    ], widths=[4.5, 11.5], caption='Bảng 3.5: Máy trạng thái nhận diện trực tiếp')
    text(d, 'Nhóm "không có rác" (mặt người, tay không, mặt bàn...) giúp camera không tự mở nắp khi chưa có rác. Delegate GPU '
            'được thử trước, lỗi hoặc ra giá trị NaN thì tự chuyển sang CPU.')

    d.add_heading('3.4.3. Đồng bộ ngoại tuyến', level=3)
    text(d, 'Bộ máy đồng bộ chạy khi mạng được khôi phục (NetInfo) và định kỳ 15 giây: lấy tối đa 50 thao tác còn lượt thử, '
            'gửi lần lượt; thành công thì xoá khỏi hàng đợi và đánh dấu sự kiện đã đồng bộ; lỗi mạng thì dừng, lỗi khác thì tăng '
            'số lần thử. Lịch sử hiển thị gộp dữ liệu máy chủ với các sự kiện chưa đồng bộ trên máy.')

    d.add_heading('3.4.4. Tự động hoá phía máy chủ', level=3)
    table(d, ['Sự kiện', 'Xử lý tự động (trigger / hàm)'], [
        ('Ngăn vượt 60%', 'create_task_when_full → auto_create_task: tạo việc "lịch ca" cho ca gần nhất (nếu chưa có việc mở)'),
        ('Ngăn vượt 80%', 'Tạo việc khẩn trong ngày, hoặc nâng việc đang có lên khẩn; báo hộ gia đình nếu là thùng nhà họ'),
        ('Việc mới chưa có người', 'auto_dispatch_task (BEFORE INSERT) → pick_collector: nhân viên đang trực, ít việc trong ngày nhất'),
        ('Việc được tạo / giao / hoàn tất / chấm điểm', 'notify_task_created, notify_task_updated → ghi bảng notifications'),
        ('Hoàn tất việc', 'empty_bins_on_done: thùng về 0%; reward_sorting_quality: +5 điểm nếu phân loại đúng'),
        ('Mỗi 5 phút (pg_cron)', 'run_automation(): (a) thùng không có nhịp tim quá ngưỡng → offline; (b) việc chưa ai nhận quá '
                                 'hạn → tự giao, việc khẩn chậm → nhắc; (c) dự báo theo tốc độ 24 giờ, đầy trong 12 giờ → lên lịch'),
    ], widths=[4.5, 11.5], caption='Bảng 3.6: Các quy tắc tự động hoá')
    text(d, 'Ứng dụng lắng nghe bảng notifications qua Realtime; khi có dòng mới, hiển thị thông báo hệ thống (expo-notifications) '
            'và tải lại dữ liệu liên quan. Kiosk và app hộ gia đình gửi nhịp tim (hàm device_heartbeat) mỗi phút khi đang kết nối thùng.')

    d.add_heading('3.4.5. Lộ trình thu gom', level=3)
    text(d, 'Đầu vào: vị trí GPS của nhân viên (nếu gần khu vực) và toạ độ các thùng đang phụ trách. Thuật toán: láng giềng gần '
            'nhất rồi cải thiện 2-opt; nếu không có GPS thì thử xuất phát từ từng thùng và chọn lộ trình ngắn nhất. Kết quả hiển thị '
            'trên bản đồ Leaflet (WebView) và mở Google Maps để chỉ đường.')

    d.add_heading('3.5. Thiết kế phần cứng và giao thức BLE', level=2)
    table(d, ['Thành phần', 'Thiết kế'], [
        ('Vi điều khiển', 'ESP32-WROOM-32 (esp32dev), PlatformIO, thư viện ESP32Servo, ArduinoJson'),
        ('Cơ cấu chấp hành', '3 servo: hữu cơ GPIO 13, vô cơ GPIO 14, tái chế GPIO 27; mở 90°, giữ 2 giây, đóng 0°'),
        ('BLE', 'Tên "SmartBin-01"; service 6e400001-…; characteristic lệnh 6e400002-… (ghi); trạng thái 6e400003-… (đọc)'),
        ('Lệnh', '{"cmd":"open","bin":"huu_co"} → phản hồi {"status":"ok"} hoặc {"status":"error","message":…}'),
        ('Dự phòng', 'MockBinController cùng giao diện BinController, bật bằng EXPO_PUBLIC_MOCK_BLE=1'),
    ], widths=[4, 12], caption='Bảng 3.7: Thiết kế phần cứng')

    d.add_heading('3.6. Thiết kế giao diện và điều hướng', level=2)
    table(d, ['Vai trò', 'Các màn hình (tab dưới)'], [
        ('Hộ gia đình', 'Trang chủ · Lịch sử · Thống kê · Cá nhân'),
        ('Kiosk', 'Cấu hình kiosk → Màn chờ → Bỏ rác (camera AI bật sẵn) → Cảm ơn (tự quay về sau 5 giây)'),
        ('Nhân viên', 'Công việc · Lộ trình · Cá nhân (+ màn Chi tiết việc)'),
        ('Quản lý', 'Thùng rác · Điều phối · Tự động · Cá nhân'),
    ], widths=[3.5, 12.5], caption='Bảng 3.8: Điều hướng theo vai trò')
    text(d, 'Giao diện dùng tông xanh lá – xanh ngọc (gradient) nhất quán, header bo góc, thẻ bo tròn; chuông thông báo trên '
            'header mọi màn hình chính.')


# ---------------------------------------------------------------- CHƯƠNG 4
def viet_chuong4(d):
    h = d.add_heading('CHƯƠNG 4: PHÂN TÍCH KẾT QUẢ', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading('4.1. Môi trường thực hiện', level=2)
    table(d, ['Hạng mục', 'Chi tiết'], [
        ('Ứng dụng', 'React Native 0.76.5 (Hermes), Expo SDK 52, Expo Router 4, TypeScript'),
        ('Thư viện chính', 'Zustand, TanStack Query, expo-sqlite, react-native-vision-camera 4, react-native-fast-tflite, '
                           'react-native-ble-plx, expo-notifications, expo-location, Leaflet (WebView)'),
        ('Máy chủ', 'Supabase (PostgreSQL, Auth, Realtime, Storage, pg_cron)'),
        ('AI', 'MobileCLIP-S0; chuyển đổi bằng onnxslim, onnx2tf; đánh giá bằng Python'),
        ('Thiết bị thử', 'Điện thoại TECNO CM6, Android 16, arm64; máy ảo Android'),
        ('Quản lý mã nguồn', REPO),
    ], widths=[3.5, 12.5], caption='Bảng 4.1: Môi trường thực hiện')

    d.add_heading('4.2. Kết quả giao diện', level=2)
    screens(d, [('h1', 'Trang chủ: camera AI tự mở nắp'), ('h9', 'Cá nhân: điểm, tác động'),
                ('h8', 'Thống kê')], 'Hình 4.1: Giao diện Hộ gia đình')
    screens(d, [('h4', 'Thông báo'), ('k1', 'Kiosk cộng đồng'), ('c1', 'Nhân viên: Công việc')],
            'Hình 4.2: Thông báo, kiosk và danh sách việc của nhân viên')
    screens(d, [('u9', 'Lộ trình tối ưu + thông báo việc mới'), ('c3', 'Nhân viên: Cá nhân'),
                ('u4', 'Quản lý: ảnh nghiệm thu')], 'Hình 4.3: Nhân viên thu gom và nghiệm thu')
    screens(d, [('a6', 'Quản lý: Thùng rác, bản đồ'), ('a2', 'Quản lý: Điều phối'), ('a3', 'Quản lý: Tự động hoá')],
            'Hình 4.4: Giao diện Quản lý')

    d.add_heading('4.3. Kết quả mô hình AI', level=2)
    text(d, 'Tập thử: 2.749 ảnh (RealWaste [5] và nguồn bổ sung, gộp 3 nhóm). Các câu mô tả chỉ được chỉnh trên tập kiểm định; '
            'tập thử chỉ chạy một lần để báo cáo.')
    table(d, ['Phiên bản', 'Dung lượng', 'Accuracy', 'Balanced acc.', 'Độ trễ CPU (PC)', 'Android'], [
        ('ONNX FP32 (gốc)', '45,5 MB', '93,16%', '94,63%', '27 ms', '—'),
        ('TFLite FP32', '45,6 MB', '93,16%', '94,63%', '82 ms', 'Dùng mặc định'),
        ('TFLite FP16', '22,9 MB', '91,92%', '93,71%', '78 ms', 'Ra NaN — loại'),
        ('TFLite INT8 dynamic', '12,1 MB', '92,91%', '93,86%', '342 ms', 'Chạy được, chậm'),
    ], widths=[3.6, 2.2, 2.2, 2.4, 2.8, 2.8], caption='Bảng 4.2: So sánh các phiên bản mô hình')
    table(d, ['Thật \\ Dự đoán', 'Hữu cơ', 'Vô cơ', 'Tái chế', 'Recall'], [
        ('Hữu cơ (326)', '320', '4', '2', '98,16%'),
        ('Vô cơ (1.080)', '3', '1.040', '37', '96,30%'),
        ('Tái chế (1.343)', '19', '123', '1.201', '89,43%'),
    ], widths=[4, 2.6, 2.6, 2.6, 2.6], caption='Bảng 4.3: Ma trận nhầm lẫn TFLite FP32 trên tập thử')
    text(d,
         'Nhận xét: mô hình zero-shot đạt 93,16%, cao hơn 3 điểm phần trăm so với mô hình MobileNetV2 nhóm tự huấn luyện trước '
         'đó (90,08% trên cùng tập) mà không cần huấn luyện. Nhầm lẫn nhiều nhất là tái chế → vô cơ (123 ảnh), chủ yếu là bao bì '
         'nhựa bẩn/nhăn — chính người dùng cũng hay nhầm. Lượng tử hoá INT8 dynamic range giảm dung lượng 3,8 lần, chỉ mất 0,25 điểm '
         'phần trăm nhưng chậm hơn vì phải giải lượng tử trọng số mỗi lần chạy; bản FP16 cho kết quả NaN trên Android do tràn số.',
         'Chế độ trực tiếp (ngưỡng 0,8, cần 2 khung liên tiếp): 86,5% khung có rác được mở ngay, trong đó 97,4% mở đúng ngăn; '
         'khung không có rác được nhận là "Không có rác" với điểm 96–100%. Trên điện thoại TECNO CM6, mỗi khung mất khoảng '
         '490 ms với GPU và 620 ms với CPU.')

    d.add_heading('4.4. Kết quả kiểm thử', level=2)
    d.add_heading('4.4.1. Kiểm thử máy chủ', level=3)
    text(d, 'Tệp supabase/tests/test_logic_server.sql đóng vai từng người dùng (đặt JWT giả lập + chuyển vai trò CSDL giống '
            'PostgREST), chạy trong một giao dịch rồi hoàn tác nên không để lại dữ liệu. Kết quả: 44/44 phép kiểm tra đạt.')
    table(d, ['Nhóm', 'Nội dung kiểm tra', 'Kết quả'], [
        ('Kiosk (anon)', 'Ghi sự kiện ẩn danh, +2% mức đầy, nhịp tim; không ghi được sự kiện mang user_id; không đọc được việc/thông báo', '6/6'),
        ('Hộ gia đình', '+1 điểm; không ghi thay người khác; chỉ đọc hồ sơ, việc, thông báo của mình; không sửa được cấu hình', '6/6'),
        ('Ngưỡng 60/80%', 'Tự tạo việc lịch ca, tự giao người ít việc, nâng khẩn không tạo trùng, báo đúng người', '7/7'),
        ('Nhân viên', 'Chỉ thấy việc của mình; không chạy được run_automation; hoàn tất → thùng 0%, +5 điểm, thông báo', '7/7'),
        ('Tắt tự giao', 'Việc chờ; nhân viên tự nhận; không tự báo cho chính mình', '5/5'),
        ('Leo thang', 'Việc bị bỏ quên tự giao; việc khẩn chậm được nhắc; không nhắc trùng', '3/3'),
        ('Thiết bị', 'Mất nhịp tim → offline + cảnh báo; nhịp tim lại → online', '3/3'),
        ('Dự báo, ca trực, quản lý', 'Dự báo theo tốc độ; không ai trực; chỉ một người trực; quyền quản lý; ràng buộc dữ liệu', '7/7'),
    ], widths=[3.3, 10.5, 2.2], caption='Bảng 4.4: Kết quả kiểm thử máy chủ')

    d.add_heading('4.4.2. Kiểm thử logic ứng dụng', level=3)
    text(d, 'Tệp scripts/test-logic.ts kiểm tra các hàm thuần: dự báo đầy (6), lộ trình và khoảng cách (6), tác động môi trường, '
            'chuỗi ngày, huy hiệu (4), tra cứu rác có/không dấu (5), trạng thái việc (1), phân loại zero-shot trên toàn bộ 180 '
            'câu mô tả (3). Kết quả: 25/25 đạt.')

    d.add_heading('4.4.3. Kiểm thử trên điện thoại thật', level=3)
    text(d, 'Kiểm thử thủ công cả 4 vai trò trên bản cài release, gồm chế độ mất mạng và thông báo realtime giữa quản lý và '
            'nhân viên. Các lỗi phát hiện và đã sửa:')
    table(d, ['Lỗi', 'Nguyên nhân', 'Cách sửa'], [
        ('Camera tự mở nắp lặp lại', 'Vật nằm yên trước camera được nhận lại sau mỗi 5 giây', 'Thêm trạng thái "chờ vật rời khung"'),
        ('Lịch sử hiện 161 lượt, máy chủ có 100', 'Gộp cả lượt đã đồng bộ nhưng đã bị xoá trên máy chủ', 'Chỉ gộp lượt chưa đồng bộ'),
        ('Bản đồ trống', 'Máy chủ tile OSM chặn yêu cầu từ WebView không có Referer', 'Chuyển sang tile Esri'),
        ('Màn Cảm ơn kẹt ở 0 giây', 'Màn tab được giữ lại, bộ đếm không đặt lại', 'Đặt lại khi màn được mở lại'),
        ('Mô phỏng 85% không tạo việc', 'Chọn ngăn đầu tiên dù ngăn đó đã đầy sẵn', 'Chọn ngăn đang dưới ngưỡng'),
        ('Lỗi giao diện', 'Chữ tràn, bàn phím che kết quả, ký tự "→" lỗi font máy TECNO', 'Sửa bố cục, tự cuộn, thay ký tự'),
    ], widths=[4.3, 6.2, 5.5], caption='Bảng 4.5: Lỗi phát hiện khi kiểm thử')

    d.add_heading('4.5. Đánh giá mức độ đạt mục tiêu', level=2)
    table(d, ['Mục tiêu (Chương 1)', 'Kết quả'], [
        ('Hỗ trợ phân loại đúng 3 nhóm', 'Đạt: chọn tay, tra cứu, AI chụp ảnh, AI camera tự mở ngăn (93,16%)'),
        ('Khuyến khích phân loại', 'Đạt: điểm, huy hiệu, chuỗi ngày, CO₂, đánh giá từ nhân viên (+5 điểm)'),
        ('Hỗ trợ nhân viên thu gom', 'Đạt: việc, lộ trình tối ưu, nghiệm thu ảnh, chấm phân loại'),
        ('Giám sát, điều phối thời gian thực', 'Đạt: bản đồ, dự báo, điều phối, lịch ca, Realtime'),
        ('Tự động hoá vận hành', 'Đạt: tự lên lịch, tự giao việc, nhắc việc, phát hiện mất kết nối, thông báo'),
        ('Hoạt động khi mất mạng', 'Đạt: SQLite + hàng đợi đồng bộ, đã kiểm thử tắt/bật mạng'),
        ('Mô hình thùng ESP32', 'Đã dựng phần cứng và firmware; đang tích hợp BLE thật (demo bằng bản giả lập)'),
    ], widths=[5.5, 10.5], caption='Bảng 4.6: Đánh giá mức độ đạt mục tiêu')


# ---------------------------------------------------------------- CHƯƠNG 5
def viet_chuong5(d):
    h = d.add_heading('CHƯƠNG 5: KẾT LUẬN', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading('5.1. Kết quả đạt được', level=2)
    bullets(d, [
        'Xây dựng hoàn chỉnh ứng dụng Android SmartBin cho 4 nhóm người dùng, chạy trên dữ liệu thật của máy chủ Supabase.',
        'Tích hợp mô hình khoa học MobileCLIP-S0 (CVPR 2024) chạy trên điện thoại, phân loại zero-shot đạt 93,16%; camera tự nhận '
        'diện rác và mở đúng ngăn không cần thao tác.',
        'Hệ thống thu gom tự vận hành phía máy chủ: tự lên lịch trước khi thùng đầy, tự giao việc cân bằng tải, tự nhắc việc trễ, '
        'phát hiện thùng mất kết nối, thông báo thời gian thực.',
        'Kiến trúc local-first cho phép dùng khi mất mạng; phân quyền an toàn bằng RLS.',
        'Bộ kiểm thử tự động (44 kiểm thử máy chủ, 25 kiểm thử logic) và kiểm thử trên thiết bị thật.',
    ])

    d.add_heading('5.2. Hạn chế', level=2)
    bullets(d, [
        'Kết nối BLE với thùng ESP32 thật chưa hoàn thiện; phần demo dùng bộ giả lập.',
        'Mức đầy được ước lượng theo số lượt bỏ rác (2%/lượt), chưa đo bằng cảm biến.',
        'Thông báo chỉ hiện khi app đang mở hoặc chạy nền; chưa có máy chủ push khi app đã bị tắt hẳn.',
        'AI còn nhầm tái chế ↔ vô cơ với bao bì nhựa bẩn; tốc độ ~0,5 giây/khung nên phải giữ rác yên khoảng 1 giây.',
        'Chỉ mới phát hành cho Android; dự báo đầy dùng mô hình tuyến tính đơn giản.',
    ])

    d.add_heading('5.3. Hướng phát triển', level=2)
    bullets(d, [
        'Hoàn thiện BLE với ESP32; gắn cảm biến siêu âm đo mức đầy thật và gửi trạng thái qua Wi-Fi/MQTT.',
        'Tích hợp Firebase Cloud Messaging / Expo Push để thông báo cả khi app đã tắt.',
        'Lượng tử hoá INT8 hoàn toàn [6] và tinh chỉnh câu mô tả / fine-tune nhẹ để tăng tốc độ và độ chính xác nhóm tái chế.',
        'Tối ưu lộ trình nhiều xe (bài toán VRP) có ràng buộc tải trọng và khung giờ.',
        'Phát hành iOS và trang web quản trị; dự báo đầy bằng mô hình chuỗi thời gian theo ngày trong tuần.',
    ])


def tai_lieu(d):
    d.add_heading('TÀI LIỆU THAM KHẢO', level=1)
    refs = [
        'P. K. A. Vasu, H. Pouransari, F. Faghri, R. Vemulapalli, and O. Tuzel, "MobileCLIP: Fast image-text models through '
        'multi-modal reinforced training," in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR), 2024, pp. 15963–15974, '
        'doi: 10.1109/CVPR52733.2024.01511.',
        'A. Radford et al., "Learning transferable visual models from natural language supervision," in Proc. 38th Int. Conf. '
        'Mach. Learn. (ICML), vol. 139, 2021, pp. 8748–8763.',
        'M. Yang and G. Thung, "Classification of trash for recyclability status," CS229 Project Report, Stanford Univ., 2016.',
        'L. Yong, L. Ma, D. Sun, and L. Du, "Application of MobileNetV2 to waste classification," PLoS ONE, vol. 18, no. 3, '
        'Art. no. e0282336, 2023, doi: 10.1371/journal.pone.0282336.',
        'S. Single, S. Iranmanesh, and R. Raad, "RealWaste: A novel real-life data set for landfill waste classification using '
        'deep learning," Information, vol. 14, no. 12, Art. no. 633, 2023, doi: 10.3390/info14120633.',
        'B. Jacob et al., "Quantization and training of neural networks for efficient integer-arithmetic-only inference," in '
        'Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR), 2018, pp. 2704–2713, doi: 10.1109/CVPR.2018.00286.',
        'I. Ranjbar, Y. Ventikos, and M. Arashpour, "Zero-shot and few-shot multimodal plastic waste classification with '
        'vision-language models," Waste Management, vol. 202, Art. no. 114815, 2025, doi: 10.1016/j.wasman.2025.114815.',
        'G. A. Croes, "A method for solving traveling-salesman problems," Operations Research, vol. 6, no. 6, pp. 791–812, 1958.',
        'Quốc hội Việt Nam, Luật Bảo vệ môi trường số 72/2020/QH14, 2020.',
        'Chính phủ Việt Nam, Nghị định 08/2022/NĐ-CP quy định chi tiết một số điều của Luật Bảo vệ môi trường, 2022.',
        'Meta Platforms, "React Native documentation." [Online]. Available: https://reactnative.dev/docs',
        'Expo, "Expo documentation." [Online]. Available: https://docs.expo.dev',
        'Supabase, "Supabase documentation: Database, Auth, Realtime, Storage." [Online]. Available: https://supabase.com/docs',
        'Google, "LiteRT (TensorFlow Lite) documentation." [Online]. Available: https://ai.google.dev/edge/litert',
    ]
    for i, r in enumerate(refs, 1):
        p = d.add_paragraph()
        p.add_run(f'[{i}] ').bold = True
        p.add_run(r)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)


def muc_luc(d):
    d.add_heading('MỤC LỤC', level=1)
    p = d.add_paragraph()
    r = p.add_run()
    for tag, txt in (('begin', None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ('separate', None), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
            r._r.append(el)
            if tag == 'separate':
                t = OxmlElement('w:t')
                t.text = '(Bấm chuột phải vào đây › Update Field để tạo mục lục)'
                r._r.append(t)
        else:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = txt
            r._r.append(it)


def bia(d):
    para(d, 'TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT TP. HỒ CHÍ MINH', bold=True, align='center', size=14)
    for _ in range(4):
        d.add_paragraph()
    para(d, 'BÁO CÁO CUỐI KỲ', bold=True, align='center', size=20)
    para(d, 'MÔN: PHÁT TRIỂN ỨNG DỤNG DI ĐỘNG', bold=True, align='center', size=15)
    d.add_paragraph()
    para(d, 'ĐỀ TÀI:', bold=True, align='center', size=14)
    para(d, 'SMARTBIN — ỨNG DỤNG QUẢN LÝ HỆ THỐNG THÙNG RÁC THÔNG MINH', bold=True, align='center', size=17)
    for _ in range(3):
        d.add_paragraph()
    table(d, ['Họ và tên', 'MSSV'], [('', ''), ('', ''), ('', '')], widths=[8, 5])
    para(d, 'GVHD: Huỳnh Hoàng Hà', bold=True, align='center')
    para(d, 'TP. Hồ Chí Minh, tháng 10 năm 2026', italic=True, align='center')


if __name__ == '__main__':
    arch = ve_kien_truc()
    chuong = {
        2: ('BaoCao_Chuong2_CoSoLyThuyet.docx', lambda d: viet_chuong2(d)),
        3: ('BaoCao_Chuong3_PhanTichThietKe.docx', lambda d: viet_chuong3(d, arch)),
        4: ('BaoCao_Chuong4_PhanTichKetQua.docx', lambda d: viet_chuong4(d)),
        5: ('BaoCao_Chuong5_KetLuan.docx', lambda d: (viet_chuong5(d), tai_lieu(d))),
    }
    for n, (name, fn) in chuong.items():
        d = base_doc()
        fn(d)
        d.save(os.path.join(OUT, name))

    d = base_doc()
    bia(d)
    page_break(d)
    muc_luc(d)
    page_break(d)
    viet_chuong1(d)
    for n, fn in ((2, lambda: viet_chuong2(d)), (3, lambda: viet_chuong3(d, arch)),
                  (4, lambda: viet_chuong4(d)), (5, lambda: viet_chuong5(d))):
        page_break(d)
        fn()
    page_break(d)
    tai_lieu(d)
    d.save(os.path.join(OUT, 'BaoCao_CuoiKy_SmartBin.docx'))
    print('Xong:', sorted(f for f in os.listdir(OUT) if f.endswith('.docx')))
